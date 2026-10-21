"""Text extraction from PDF / DOCX / TXT (with optional OCR for scanned PDFs)."""
from __future__ import annotations

import io
import os
from typing import List, Tuple

MAX_PAGES = 400
MAX_CHARS = 800_000


class ExtractionError(Exception):
    pass


def _pdf_pdfplumber(data: bytes) -> Tuple[List[str], int]:
    import pdfplumber

    pages: List[str] = []
    with pdfplumber.open(io.BytesIO(data)) as pdf:
        n = len(pdf.pages)
        for page in pdf.pages[:MAX_PAGES]:
            try:
                pages.append(page.extract_text(x_tolerance=2, y_tolerance=3) or "")
            except Exception:
                pages.append("")
    return pages, n


def _pdf_pypdf(data: bytes) -> Tuple[List[str], int]:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data))
    if reader.is_encrypted:
        try:
            reader.decrypt("")
        except Exception as e:
            raise ExtractionError("The PDF is password-protected.") from e
    pages = []
    for p in reader.pages[:MAX_PAGES]:
        try:
            pages.append(p.extract_text() or "")
        except Exception:
            pages.append("")
    return pages, len(reader.pages)


def _pdf_ocr(data: bytes, warnings: List[str]) -> List[str]:
    try:
        import pdfplumber
        import pytesseract  # type: ignore
    except Exception:
        warnings.append(
            "This PDF looks scanned (no selectable text). Install OCR support to read it: "
            "`pip install pytesseract` plus the Tesseract program (see README)."
        )
        return []
    pages = []
    try:
        with pdfplumber.open(io.BytesIO(data)) as pdf:
            for page in pdf.pages[:60]:
                img = page.to_image(resolution=200).original
                pages.append(pytesseract.image_to_string(img))
        warnings.append("Scanned PDF read with OCR; accuracy depends on scan quality.")
    except Exception as e:
        warnings.append(f"OCR failed: {type(e).__name__}: {str(e)[:120]}")
    return pages


def extract_pdf(data: bytes, enable_ocr: bool = True) -> Tuple[str, List[str], int]:
    warnings: List[str] = []
    pages: List[str] = []
    n_pages = 0
    try:
        pages, n_pages = _pdf_pdfplumber(data)
    except Exception as e1:
        try:
            pages, n_pages = _pdf_pypdf(data)
            warnings.append("pdfplumber could not read this PDF; used pypdf instead.")
        except ExtractionError:
            raise
        except Exception as e2:
            raise ExtractionError(f"Could not read this PDF ({type(e1).__name__}; {type(e2).__name__}).") from e2
    text = "\n\n".join(pages).strip()
    if n_pages and len(text) < 40 * max(1, min(n_pages, 5)):
        if enable_ocr:
            ocr_pages = _pdf_ocr(data, warnings)
            if ocr_pages:
                text = "\n\n".join(ocr_pages).strip()
        else:
            warnings.append("No selectable text found (scanned PDF?). Enable OCR to read it.")
    if n_pages > MAX_PAGES:
        warnings.append(f"Only the first {MAX_PAGES} pages were analysed.")
    return text, warnings, n_pages


def extract_docx(data: bytes) -> Tuple[str, List[str], int]:
    try:
        import docx  # python-docx
    except Exception as e:
        raise ExtractionError("python-docx is not installed (pip install python-docx).") from e
    try:
        d = docx.Document(io.BytesIO(data))
    except Exception as e:
        raise ExtractionError(f"Could not read this Word file: {e}") from e
    parts = [p.text for p in d.paragraphs]
    for table in d.tables:
        for row in table.rows:
            cells = []
            for c in row.cells:
                if c.text not in cells:
                    cells.append(c.text)
            parts.append(" | ".join(cells))
    return "\n".join(parts), [], 0


def extract_txt(data: bytes) -> Tuple[str, List[str], int]:
    for enc in ("utf-8-sig", "utf-16", "cp1252", "latin-1"):
        try:
            return data.decode(enc), [], 0
        except Exception:
            continue
    return data.decode("utf-8", errors="ignore"), [], 0


def extract_text(data: bytes, filename: str = "document.txt", enable_ocr: bool = True) -> Tuple[str, List[str], int]:
    """Return (text, warnings, page_count)."""
    if not data:
        raise ExtractionError("The file is empty.")
    ext = os.path.splitext(filename.lower())[1]
    if ext == ".pdf" or data[:5] == b"%PDF-":
        text, warns, pages = extract_pdf(data, enable_ocr)
    elif ext == ".docx" or data[:2] == b"PK":
        text, warns, pages = extract_docx(data)
    elif ext in (".txt", ".md", ".text", ""):
        text, warns, pages = extract_txt(data)
    else:
        raise ExtractionError(f"Unsupported file type '{ext}'. Use PDF, DOCX or TXT.")
    if len(text) > MAX_CHARS:
        text = text[:MAX_CHARS]
        warns.append("The document is very long; only the first part was analysed.")
    return text, warns, pages
